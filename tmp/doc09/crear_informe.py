from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
D=Document(); sec=D.sections[0];sec.page_height=Cm(29.7);sec.page_width=Cm(21);sec.top_margin=sec.bottom_margin=Cm(2);sec.left_margin=sec.right_margin=Cm(2.2)
for n in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 st=D.styles[n];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
D.styles['Normal'].font.size=Pt(11);D.styles['Normal'].paragraph_format.space_after=Pt(7)
D.styles['Normal'].paragraph_format.line_spacing=1.12
D.styles['Title'].font.size=Pt(23)
D.styles['Heading 1'].font.size=Pt(16)
D.styles['Heading 2'].font.size=Pt(12)
D.core_properties.title='Informe de capacidades nativas en PharmaMobil';D.core_properties.subject='Sesión 09 expect actual y compartir productos';D.core_properties.author=''
def p(t):D.add_paragraph(t)
def h(t):D.add_heading(t,level=1)
def sub(t):D.add_heading(t,level=2)
def bullet(t):D.add_paragraph(t,style='List Bullet')
def code(t):
 r=D.add_paragraph().add_run(t);r.font.name='Consolas';r.font.size=Pt(9)
def page():D.add_page_break()
def table(headers,rows,widths=None):
 t=D.add_table(rows=1, cols=len(headers));t.autofit=False
 for i,x in enumerate(headers):t.rows[0].cells[i].text=x
 for row in rows:
  cells=t.add_row().cells
  for i,x in enumerate(row):cells[i].text=x
 for ri,row in enumerate(t.rows):
  for i,c in enumerate(row.cells):
   if widths:c.width=Cm(widths[i])
   pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   pr.append(b)
   margins=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margins.append(e)
   pr.append(margins)
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7EDF2');pr.append(sh)
   for pp in c.paragraphs:
    pp.paragraph_format.space_after=Pt(3)
    for r in pp.runs:r.font.size=Pt(10);r.bold=ri==0
  trpr=row._tr.get_or_add_trPr();ns=OxmlElement('w:cantSplit');trpr.append(ns)
 t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
def evidence(title,description):
 sub(title);p(description)
 pp=D.add_paragraph('[Insertar captura aquí]');pp.paragraph_format.space_after=Pt(55);pp.runs[0].italic=True
 p('Dispositivo y versión del sistema: ____________________')
 p('Resultado observado: ______________________________________________')
D.add_paragraph('Informe de capacidades nativas en PharmaMobil',style='Title')
D.add_paragraph('Práctica de laboratorio 09',style='Subtitle')
p('Desarrollo de Aplicaciones Móviles | Ciclo VI | Semestre 2026 2')
p('Fecha de la práctica: 6 de octubre de 2026')
p('Integrante 1: _________________________________________________')
p('Integrante 2: _________________________________________________')
p('Docente: Mg. Reyna Barreto Benjamin David')
h('1 Propósito y alcance')
p('La práctica incorpora dos capacidades nativas a PharmaMobil: mostrar los precios de los productos en soles y compartir sus datos mediante el sistema operativo. El código común conserva la interfaz y la coordinación de las acciones; Android e iOS aportan implementaciones específicas.')
p('El formato se declara con expect y se implementa con actual. Compartir se define mediante un contrato del dominio y se resuelve con Koin. La compilación de Android y las pruebas automatizadas finalizaron correctamente. La ejecución de compartir y las evidencias visuales deben completarse en los dispositivos; la validación de iOS requiere macOS y Xcode.')
h('2 Objetivos')
bullet('Implementar formatearSoles en commonMain, androidMain e iosMain.')
bullet('Preparar el precio visible en la capa de presentación sin cambiar el valor numérico del dominio.')
bullet('Abrir el selector de compartir desde el detalle del producto mediante una interfaz común.')
bullet('Documentar la organización por capas, las diferencias entre plataformas y las evidencias de funcionamiento.')
page();h('3 Organización del código por capas')
p('Las rutas siguientes son relativas a shared/src. Dentro de cada source set, los archivos Kotlin se encuentran bajo kotlin/pe/edu/upeu/pharmamobile/.')
table(['Capa o ubicación','Responsabilidad','Motivo'],[
('commonMain/domain','Producto, Compartidor y texto compartido','Define los datos y el contrato sin importar Android ni UIKit.'),
('commonMain/presentation','ProductoUi, pantallas y ViewModels','Prepara los datos visibles y coordina las acciones del usuario.'),
('commonMain/data','API y repositorio REST existentes','Consulta y persiste productos en PharmaSoft.'),
('commonMain/platform','Declaración expect del formato','Expone una firma común para la utilidad de moneda.'),
('androidMain/platform','Formato y compartir para Android','Utiliza bibliotecas Java y APIs nativas de Android.'),
('iosMain/platform','Formato y compartir para iOS','Utiliza Foundation y UIKit.'),
('di en cada source set','Módulos de Koin','Conecta cada contrato con su implementación concreta.')],[4,6,6.5])
sub('Separación entre precio y presentación')
p('Producto mantiene precio como Double para cálculos y operaciones. ProductoUi contiene el producto original y precioFormateado como String. La conversión toUi se realiza al preparar el estado en ProductoViewModel, antes de dibujar el listado y el detalle.')
sub('Recorrido de compartir')
p('Productos → Ver detalle → Compartir → DetalleProductoViewModel → Compartidor → implementación nativa. El ViewModel prepara nombre, precio y stock; Koin proporciona CompartidorAndroid o CompartidorIos según la plataforma.')
p('La función comoTextoParaCompartir se conserva en domain/usecase siguiendo la guía. Invoca la utilidad común de moneda para preparar el mensaje, aunque ese formato representa una decisión de presentación.')
page();h('4 Implementación de las capacidades')
sub('Formato de moneda con expect y actual')
code('commonMain/platform/Formato.kt\nexpect fun formatearSoles(valor: Double): String')
p('Android implementa la función en platform/Formato.android.kt mediante NumberFormat y Locale es PE. iOS la implementa en platform/Formato.ios.kt mediante NSNumberFormatter, NSLocale es_PE y moneda PEN. Las declaraciones comparten paquete, nombre, parámetros y retorno.')
p('ProductoUi.kt llama a formatearSoles dentro de toUi. FaseProductos.ConProductos almacena una lista de ProductoUi; los filtros, la edición y la eliminación siguen accediendo al producto original.')
sub('Contrato y texto compartido')
code('commonMain/domain/platform/Compartidor.kt\ninterface Compartidor {\n    fun compartir(texto: String)\n}')
p('domain/usecase/TextoParaCompartir.kt define Producto.comoTextoParaCompartir. El mensaje incluye el nombre, el precio formateado y el stock. Ejemplo de contenido: Paracetamol — S/ 5.50 — Stock: 20. La representación exacta del símbolo y los espacios depende del formateador nativo.')
sub('Implementaciones nativas')
p('CompartidorAndroid utiliza Intent.ACTION_SEND, tipo text/plain y EXTRA_TEXT. Intent.createChooser permite elegir la aplicación receptora. FLAG_ACTIVITY_NEW_TASK permite abrir el selector desde el contexto de la aplicación proporcionado por Koin.')
p('CompartidorIos crea UIActivityViewController, obtiene una ventana activa y presenta la hoja desde el controlador visible. Configura sourceView y sourceRect para el popover de iPad. Su compilación y comportamiento están pendientes de verificar en macOS.')
sub('Inyección y pantalla de detalle')
p('PlatformModule.android.kt y PlatformModule.ios.kt registran Compartidor y conservan sus motores HTTP OkHttp y Darwin. AppModule.kt registra DetalleProductoViewModel e incluye initKoinIos como entrada para Swift.')
p('presentation/detalle/DetalleProductoScreen.kt presenta un diálogo con nombre, precio, stock y estado, además de Compartir y Cerrar. ProductoScreen abre ese diálogo mediante Ver detalle. Las pantallas no importan APIs nativas.')
page();h('5 Diferencias entre Android e iOS')
table(['Aspecto','Android','iOS'],[
('Formato','NumberFormat y Locale','NSNumberFormatter y NSLocale'),
('Compartir','ACTION_SEND y selector de aplicaciones','UIActivityViewController y hoja de compartir'),
('Dependencia nativa','Context de la aplicación','Ventana activa y controlador visible'),
('Presentación','Nueva actividad desde Context','Presentación sobre un controlador; popover en iPad'),
('Entorno de prueba','Android Studio en Windows','macOS con Xcode'),
('Opciones disponibles','Dependen de las aplicaciones instaladas','Dependen de las aplicaciones y actividades disponibles')],[3.5,6.5,6.5])
sub('Interoperabilidad con Swift')
p('El módulo shared genera el framework Shared. iOSApp.swift importa ese framework e inicializa Koin mediante AppModuleKt.initKoinIos(). ContentView utiliza MainViewControllerKt.MainViewController() para abrir la interfaz Compose. El sufijo Kt identifica las funciones Kotlin de nivel superior expuestas a Swift.')
h('6 Validación y resultados')
p('Se ejecutó la siguiente verificación en Windows:')
code('gradlew.bat :androidApp:assembleDebug\n            :shared:testAndroidHostTest --console=plain')
p('Resultado registrado: BUILD SUCCESSFUL. Se generó la compilación debug de Android y se ejecutaron las pruebas del módulo shared. Se adaptó la prueba del listado para comprobar los productos originales contenidos en ProductoUi.')
p('La revisión del código común no encontró importaciones de android ni de platform.UIKit. Esto conserva la independencia de las pantallas respecto de cada plataforma.')
sub('Comprobación manual pendiente')
bullet('Ejecutar Android, abrir Productos y comprobar el precio visible.')
bullet('Abrir Ver detalle y Compartir; seleccionar una aplicación y comprobar el texto recibido.')
bullet('Compilar y ejecutar iOS en una Mac; comprobar formato, hoja de compartir y comportamiento en iPad si está disponible.')
bullet('Insertar las capturas de las páginas siguientes y completar el dispositivo y resultado observado.')
page();h('7 Evidencias del formato y Android')
evidence('Figura 1 Listado con precios en Android','Capturar el listado de productos con precios formateados en soles.')
evidence('Figura 2 Listado con precios en iOS','Capturar el mismo listado en iOS. Registrar diferencias visibles en símbolo, separadores o espacios.')
evidence('Figura 3 Selector de compartir en Android','Abrir Ver detalle y Compartir. Capturar el selector del sistema con las aplicaciones disponibles.')
page();h('8 Evidencias de iOS y control inicial')
evidence('Figura 4 Hoja de compartir en iOS','Capturar la hoja nativa de compartir con el texto del producto.')
evidence('Figura 5 Error por falta de actual','Insertar la captura del punto de control 1, obtenida después de declarar expect y antes de crear las implementaciones actual. Esta evidencia aún no está incorporada.')
h('9 Datos de entrega')
p('Rama del integrante 1: ______________________________________________')
p('Enlace a la rama: __________________________________________________')
p('Tres commits propios: ______________________________________________')
p('Rama del integrante 2: ______________________________________________')
p('Enlace a la rama y tres commits: ______________________________________')
p('El README incluye la sección Capacidades nativas. Según la guía, la actividad autónoma se entrega hasta el 12 de octubre de 2026 a las 23:59.')
h('10 Conclusiones y referencia')
p('Las capacidades incorporadas permiten conservar una interfaz común y delegar las operaciones nativas a cada plataforma. Android cuenta con compilación y pruebas automatizadas satisfactorias. La comprobación funcional en dispositivos y la validación de iOS completarán las evidencias de la práctica.')
p('Referencia: Guía Práctica de Laboratorio N.º 09, Capacidades nativas con expect/actual en PharmaMobil, Desarrollo de Aplicaciones Móviles, 6 de octubre de 2026. Archivos del proyecto PharmaMobil revisados para esta implementación.')
D.save('output/doc/Informe_Sesion09_PharmaMobil.docx')
print('Word creado')
