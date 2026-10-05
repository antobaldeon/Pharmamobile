from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
import json,subprocess
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]; TMP=ROOT/'tmp/s07'
OUT=ROOT/'output/documents';OUT.mkdir(parents=True,exist_ok=True)
BASE=ROOT/'shared/src/commonMain/kotlin/pe/edu/upeu/pharmamobile'
data=json.loads((TMP/'productos.json').read_text(encoding='utf-8'))
queries=json.loads((TMP/'consultas.json').read_text(encoding='utf-8'))
commit=subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip()
doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.9);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
 s=doc.styles[name];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
 s.paragraph_format.space_after=Pt(7)
 for border in s.element.xpath('.//w:pBdr'):
  border.getparent().remove(border)
doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.line_spacing=1.12
doc.styles['Title'].font.size=Pt(27)
doc.styles['Heading 1'].font.size=Pt(17);doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Caption'].font.size=Pt(9)
code_style=doc.styles.add_style('Código Kotlin',1);code_style.font.name='Consolas';code_style.font.size=Pt(8)
code_style.paragraph_format.space_after=Pt(0);code_style.paragraph_format.line_spacing=1
code_style.paragraph_format.keep_with_next=False
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('PharmaMobil  |  Sesión 07  |  ');r.font.size=Pt(8)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
doc.core_properties.author='Tania Antonella Baldeon Ramirez'
doc.core_properties.title='Documentación de endpoints DTO y pruebas de conexión de PharmaMobil'
def p(text='',style=None):return doc.add_paragraph(text,style)
def h(text,level=1):doc.add_heading(text,level)
def field(label,value):
 q=p();q.add_run(label+' ').bold=True;q.add_run(value);return q
def page():doc.add_page_break()
def code(text):
 for line in text.strip().splitlines():p(line,'Código Kotlin')
 p()
def table(headers,rows,widths):
 t=doc.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for i,w in enumerate(widths):t.columns[i].width=Cm(w)
 for i,v in enumerate(headers):t.rows[0].cells[i].text=v
 for row in rows:
  cells=t.add_row().cells
  for i,v in enumerate(row):cells[i].text=str(v)
 for ridx,row in enumerate(t.rows):
  trPr=row._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');trPr.append(no)
  if ridx==0:
   repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
  for i,c in enumerate(row.cells):
   c.width=Cm(widths[i]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9');b.append(edge)
   pr.append(b);m=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    x=OxmlElement('w:'+side);x.set(qn('w:w'),'90');x.set(qn('w:type'),'dxa');m.append(x)
   pr.append(m)
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCEAF2' if ridx==0 else ('F5F8FA' if ridx%2==0 else 'FFFFFF'));pr.append(shade)
   for q in c.paragraphs:
    q.paragraph_format.space_after=Pt(3);q.paragraph_format.line_spacing=1.05
    for r in q.runs:r.font.size=Pt(9);r.bold=ridx==0
 p()
 return t
def image(path,width,caption):
 q=p();q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.add_run().add_picture(str(path),width=Cm(width))
 q.paragraph_format.space_after=Pt(3)
 p(caption,'Caption')

p('UNIVERSIDAD PERUANA UNIÓN','Heading 2')
p('Facultad de Ingeniería y Arquitectura\nEscuela Profesional de Ingeniería de Sistemas')
p('Desarrollo de Aplicaciones Móviles\nCiclo VI  |  Semestre 2026 2')
p();p('Documentación de endpoints DTO y pruebas de conexión','Title')
p('Actividad Autónoma 07\nProyecto PharmaMobil','Subtitle')
p();field('Estudiante','Tania Antonella Baldeon Ramirez')
field('Fecha','1 de octubre de 2026')
field('Modalidad','Individual sobre el proyecto de pareja')
field('Repositorio','https://github.com/antobaldeon/Pharmamobile')
field('Rama','feature/ktor-client')
h('Propósito del informe',2)
p('Documentar el contrato REST de productos, la correspondencia entre las respuestas JSON y los DTO de Kotlin, y el comportamiento de la conexión de PharmaMobil. Esta información orienta la implementación del CRUD de la Unidad 2.')
p('La aplicación Android registra una respuesta HTTP 200 del listado de productos. El backend devuelve cuatro registros y el contrato OpenAPI confirma las cinco operaciones CRUD. La bitácora diferencia los resultados comprobados de los escenarios que requieren completar su ejecución móvil.')
h('Contenido',2)
p('1  Catálogo de endpoints\n2  Diccionario de DTO\n3  JSON real y transformación al dominio\n4  Bitácora de pruebas de conexión\n5  Evidencias y repositorio\n6  Conclusiones y lista de cotejo')

page();h('1 Catálogo de endpoints')
field('URL base en Android','http://10.0.2.2:8080/')
field('URL local del backend','http://127.0.0.1:8080/')
field('Versión de las rutas','v1')
field('Contrato consultado','http://127.0.0.1:8080/v3/api-docs')
p('OpenAPI muestra info.version = v0, mientras que las rutas del servicio usan /api/v1. Se conserva esta distinción para evitar confundir la versión de la documentación con el prefijo de los endpoints.')
table(['Método y ruta','Parámetros o cuerpo','Respuesta','Errores'],[
['GET\n/api/v1/productos','Query: pagina, tamanio, ordenarPor, direccion','200\nPágina de productos','No declarados en OpenAPI'],
['GET\n/api/v1/productos/{id}','Ruta: id de tipo Long','200\nProductoResponseDTO','400 ID inválido; 404 inexistente, comprobados'],
['POST\n/api/v1/productos','JSON: nombre, precio, stock, estado, categoriaId','200 según OpenAPI\nProductoResponseDTO','No declarados en OpenAPI'],
['PUT\n/api/v1/productos/{id}','Ruta: id\nJSON: mismo cuerpo de creación','200 según OpenAPI\nProductoResponseDTO','No declarados en OpenAPI'],
['DELETE\n/api/v1/productos/{id}','Ruta: id de tipo Long','200 según OpenAPI\nSin esquema de cuerpo','No declarados en OpenAPI']], [4.5,5.2,3.5,3.8])
h('Parámetros de listado y validación',2)
p('El cliente envía pagina desde 0, tamanio=20, ordenarPor=id y direccion=asc. El repositorio incrementa pagina y acumula contenido hasta recibir ultima=true. Estos valores coinciden con los valores por defecto publicados por el backend.')
p('El cuerpo de creación y actualización exige nombre de 3 a 150 caracteres, precio mínimo 0.01, stock mínimo 0, estado booleano y categoriaId. Estos DTO de solicitud pertenecen al backend y todavía no están declarados en el cliente móvil.')
p('El listado y las consultas por ID fueron ejecutados como consultas de lectura. POST, PUT y DELETE están confirmados en OpenAPI, pero su ejecución y sus códigos de error requieren pruebas adicionales. La aplicación móvil implementa el listado; registrar aún lanza UnsupportedOperationException.')

page();h('2 Diccionario de DTO')
p('Los DTO del cliente se encuentran en data/remote/dto/ProductosDto.kt. Todos sus campos son obligatorios para la deserialización, no anulables y sin valores por defecto.')
table(['DTO y campo JSON','Tipo Kotlin','Obligatorio','Valor por defecto','Campo en dominio'],[
['ProductoDto.id','Long','Sí','Ninguno','Producto.id'],
['ProductoDto.nombre','String','Sí','Ninguno','Producto.nombre'],
['ProductoDto.precio','Double','Sí','Ninguno','Producto.precio'],
['ProductoDto.stock','Int','Sí','Ninguno','Producto.stock'],
['ProductoDto.estado','Boolean','Sí','Ninguno','Producto.activo'],
['PaginaProductosDto.contenido','List<ProductoDto>','Sí','Ninguno','List<Producto> mediante toDomain()'],
['PaginaProductosDto.ultima','Boolean','Sí','Ninguno','Control de paginación']], [4.7,3.5,2.1,2.4,4.3])
h('Declaración Kotlin',2)
code((BASE/'data/remote/dto/ProductosDto.kt').read_text(encoding='utf-8'))
h('Campos adicionales de la respuesta',2)
p('El backend también devuelve categoriaId, categoriaNombre, fechaCreacion y fechaModificacion en cada producto, además de pagina, tamanio, totalElementos y totalPaginas en la página. El cliente no los declara. ignoreUnknownKeys=true permite consumir esta respuesta sin modelar esos campos adicionales.')

page();h('3 JSON real y transformación al dominio')
p('Fragmento de la respuesta real de GET /api/v1/productos. Se conserva el primer producto y los metadatos; los otros tres productos se omiten únicamente para hacer legible el fragmento.')
fragment={**data,'contenido':data['contenido'][:1]}
code(json.dumps(fragment,ensure_ascii=False,indent=2))
p('El servidor devolvió cuatro productos: Paracetamol, Quitadol, Dologran y Amoxicilina. totalElementos=4, totalPaginas=1 y ultima=true indican que el listado queda completo en una sola petición.')
h('Correspondencia con el modelo',2)
code((BASE/'data/mapper/ProductoMapper.kt').read_text(encoding='utf-8'))
p('La transformación mantiene id, nombre, precio y stock, y convierte estado en activo. El DTO de página se usa en la capa de datos; el dominio recibe una lista de Producto.')

page();h('4 Bitácora de pruebas de conexión')
h('P01 Respuesta exitosa',2)
field('Fecha y hora','1 de octubre de 2026 a las 23:12:20 en Lima, equivalente a 02/10/2026 04:12:20 GMT en las cabeceras del servidor')
field('Plataforma','Android en emulador Pixel 9a')
field('Petición','GET /api/v1/productos?pagina=0&tamanio=20&ordenarPor=id&direccion=asc')
field('Pasos','Ejecutar la aplicación con el backend disponible y solicitar el listado de productos. Consultar el registro Ktor de la petición y su respuesta.')
field('Resultado esperado','HTTP 200 y productos mostrados en la interfaz.')
field('Resultado observado','Ktor registró RESPONSE: 200 y Content-Type: application/json. La consulta directa al backend devolvió cuatro productos.')
field('Conclusión','La petición de listado alcanza el servidor y recibe una respuesta satisfactoria. La captura siguiente documenta el estado visible de la aplicación.')
shot=TMP/'productos_android.png'
image(shot,5.4,'Figura 1  Aplicación PharmaMobil ejecutada en Android. La pantalla capturada debe contrastarse con el listado para acreditar también su renderizado.')

page();h('4 Bitácora de pruebas de conexión continuación')
h('P02 Recurso inexistente',2)
field('Fecha y hora','1 de octubre de 2026 a las 23:15:07 en Lima')
field('Plataforma','Cliente HTTP en Windows; comprobación del backend')
field('Pasos','Consultar GET /api/v1/productos/9223372036854775807, identificador inexistente.')
field('Resultado esperado','HTTP 404. En la app, ClientRequestException capturada y un mensaje controlado.')
field('Resultado observado','El backend respondió HTTP 404 con el mensaje “Producto no encontrado con id: 9223372036854775807”.')
field('Conclusión','El código 404 está comprobado en el servidor. Falta ejecutar el escenario mediante Ktor y capturar el mensaje de la aplicación, porque el cliente actual no tiene consulta por ID.')
h('P03 Sin conexión',2)
field('Pasos pendientes','Activar modo avión antes de solicitar el listado y registrar la pantalla de error; restaurar la conectividad al terminar.')
field('Resultado esperado','Excepción de entrada/salida capturada y mensaje visible al usuario.')
field('Estado','No ejecutado con modo avión. El log previo contiene ConnectException cuando el backend no estaba disponible; eso corresponde a servidor inaccesible y no acredita ausencia de internet.')
h('P04 Tiempo de espera agotado',2)
field('Pasos pendientes','Cambiar temporalmente requestTimeoutMillis de 15000 a 1, solicitar productos, capturar el resultado y restaurar la configuración.')
field('Resultado esperado','HttpRequestTimeoutException controlada sin congelar la interfaz.')
field('Estado','Pendiente de ejecución y captura móvil.')
h('P05 Campo desconocido en JSON',2)
field('Resultado observado parcial','La respuesta real contiene campos no declarados en los DTO y Ktor registra HTTP 200 con ignoreUnknownKeys=true. El log de cabeceras no acredita por sí solo el resultado final de deserialización.')
field('Comparación pendiente','Repetir la consulta con ignoreUnknownKeys=false, registrar el error de serialización y restaurar true.')
field('Resultado esperado','Con true se ignoran los campos extra; con false la deserialización falla.')
p('Para cerrar P02 a P05, completar fecha y hora de ejecución móvil, pasos exactos, resultado observado, captura y conclusión. El ViewModel muestra error.message o el texto “No se pudieron cargar los productos”; su contenido real debe registrarse en cada prueba.')

page();h('5 Evidencias y repositorio')
h('Registro real de Ktor',2)
p('Extracto del log de Android. La cabecera Date usa GMT; 04:12:20 del 2 de octubre equivale a 23:12:20 del 1 de octubre en Lima.')
log=(TMP/'ktor.log').read_text(encoding='utf-8')
lines=[x for x in log.splitlines() if '04:12:19.' in x or '04:12:20.' in x]
clean=[]
for x in lines:
 line=x.split('I System.out: ',1)[-1]
 if len(line)>90:
  if 'http://' in line:
   a,b=line.split('http://',1);clean.extend([a.strip(),'http://'+b[:75],b[75:]] if len(b)>75 else [a.strip(),'http://'+b])
  else:clean.extend([line[:90],line[90:]])
 else:clean.append(line)
code('\n'.join(clean))
p('El cliente utiliza LogLevel.HEADERS. El extracto acredita método, URL, cabeceras y estado, pero no contiene el cuerpo completo. Para la evidencia solicitada de petición y respuesta completas, capturar una ejecución con el registro de cuerpos habilitado.')
h('Repositorio y commit',2)
field('Repositorio','https://github.com/antobaldeon/Pharmamobile')
field('Rama local','feature/ktor-client')
field('Commit del código documentado',commit)
p('https://github.com/antobaldeon/Pharmamobile/commit/'+commit)
p('El README incluye la sección Conectividad REST. El informe y los cambios locales del README deben incorporarse a un nuevo commit si se desea que formen parte de la entrega; el hash anterior corresponde al código analizado.')
h('Ejecución en iOS',2)
p('El proyecto contiene iosApp y el motor Darwin en iosMain. La evidencia de ejecución en simulador iOS está pendiente: requiere macOS y Xcode. Una configuración de pruebas iOS en Android Studio no sustituye una captura de la aplicación ejecutándose.')

page();h('6 Conclusiones y lista de cotejo')
p('PharmaMobil tiene un cliente Ktor configurado para consumir el listado paginado del backend y convertir sus DTO al modelo de dominio. La respuesta HTTP 200 registrada en Android y el JSON real verifican que el servicio entrega los productos. El contrato OpenAPI permite documentar las operaciones necesarias para ampliar el cliente al CRUD.')
p('El manejo de errores del listado se concentra en el ViewModel mediante runCatching y un estado de error. Para completar la actividad, se deben acreditar con ejecución móvil los errores 404, ausencia de conexión, timeout y el contraste de campos desconocidos, junto con sus mensajes y capturas.')
table(['Requisito de entrega','Situación'],[
['Portada con nombre, sesión, fecha y repositorio','Incluido'],
['Catálogo de cinco endpoints','Incluido y contrastado con OpenAPI; faltan pruebas de escritura'],
['Diccionario de todos los DTO móviles','Incluido'],
['JSON real y código Kotlin correspondiente','Incluidos'],
['P01 con registro y captura Android','HTTP 200 comprobado; revisar captura de lista'],
['P02 a P05 con resultados y capturas móviles','Pendientes de completar'],
['Registro de Ktor con cuerpos completos','Pendiente; registro actual de cabeceras'],
['Captura de simulador iOS','Pendiente'],
['README Conectividad REST','Actualizado localmente'],
['Commit de la entrega accesible','Código identificado; cambios del informe por publicar']], [9.3,7.7])
h('Referencias',2)
p('Universidad Peruana Unión. Actividad Autónoma N.º 07. Documentación de endpoints, DTO y pruebas de conexión. Desarrollo de Aplicaciones Móviles, semestre 2026-2.')
p('PharmaMobil. Código fuente de HttpClientFactory, ProductoApi, ProductosDto, ProductoMapper, ProductoRepositoryImpl y ProductoViewModel. Rama feature/ktor-client.')
p('Backend local PharmaMobil. Contrato OpenAPI /v3/api-docs y consultas GET /api/v1/productos y /api/v1/productos/{id}, realizadas el 1 de octubre de 2026.')
dest=OUT/'S07_ActividadAutonoma_Baldeon_Ramirez.docx';doc.save(dest);print(dest)
